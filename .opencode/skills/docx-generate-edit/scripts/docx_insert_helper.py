import re
import unicodedata
from dataclasses import dataclass
from typing import Iterable, Optional, Sequence

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


@dataclass(frozen=True)
class ParagraphBlock:
    style: str
    text: str


@dataclass(frozen=True)
class InsertOptions:
    page_break_before_first_block: bool = False
    skip_if_text_exists: Optional[str] = None


@dataclass(frozen=True)
class InsertPosition:
    mode: str = "end"
    anchor_text: Optional[str] = None
    paragraph_index: Optional[int] = None


def normalize_style_name(name: str) -> str:
    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]", "", ascii_name.lower())


def get_paragraph_style_names(document: Document) -> list[str]:
    return [style.name for style in document.styles if style.type == WD_STYLE_TYPE.PARAGRAPH]


def get_paragraph_styles(document: Document) -> list:
    return [style for style in document.styles if style.type == WD_STYLE_TYPE.PARAGRAPH]


def resolve_paragraph_style(document: Document, requested_style: str) -> str:
    available_styles = get_paragraph_styles(document)
    available_style_names = [style.name for style in available_styles]
    normalized_requested = normalize_style_name(requested_style)

    for style in available_styles:
        if normalize_style_name(style.name) == normalized_requested:
            return style.name

    for style in available_styles:
        if normalize_style_name(style.style_id) == normalized_requested:
            return style.name

    raise ValueError(
        f"Unable to resolve paragraph style '{requested_style}'. "
        f"Available paragraph styles: {', '.join(sorted(available_style_names))}"
    )


def find_paragraph_by_text(document: Document, target_text: str) -> Paragraph:
    normalized_target = target_text.strip()
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == normalized_target:
            return paragraph

    raise ValueError(f"Unable to find paragraph with text: {target_text!r}")


def insert_blocks(
    docx_path: str,
    blocks: Iterable[ParagraphBlock],
    *,
    position: Optional[InsertPosition] = None,
    options: Optional[InsertOptions] = None,
) -> None:
    """Insert styled paragraphs into a .docx file.

    Supported positions:
    - end
    - start
    - index
    - before_text
    - after_text
    """
    document = Document(docx_path)
    config = options or InsertOptions()
    target = position or InsertPosition()
    block_list: Sequence[ParagraphBlock] = tuple(blocks)

    if not block_list:
        return

    if config.skip_if_text_exists:
        existing_texts = {paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()}
        if config.skip_if_text_exists in existing_texts:
            return

    if target.mode == "end":
        _insert_at_end(document, block_list, config)
    elif target.mode == "start":
        _insert_at_start(document, block_list, config)
    elif target.mode == "index":
        _insert_at_index(document, block_list, target, config)
    elif target.mode == "before_text":
        _insert_before_text(document, block_list, target, config)
    elif target.mode == "after_text":
        _insert_after_text(document, block_list, target, config)
    else:
        raise ValueError(
            "Unsupported insertion mode. Use one of: end, start, index, before_text, after_text"
        )

    document.save(docx_path)


def _insert_at_end(document: Document, blocks: Sequence[ParagraphBlock], options: InsertOptions) -> None:
    if options.page_break_before_first_block:
        page_break_paragraph = document.add_paragraph(style=resolve_paragraph_style(document, "Normal"))
        page_break_paragraph.add_run().add_break(WD_BREAK.PAGE)

    for block in blocks:
        paragraph = document.add_paragraph(style=resolve_paragraph_style(document, block.style))
        paragraph.add_run(block.text)


def _insert_at_start(document: Document, blocks: Sequence[ParagraphBlock], options: InsertOptions) -> None:
    if document.paragraphs:
        anchor = document.paragraphs[0]
        _insert_before_anchor(document, anchor, blocks, options)
        return

    _insert_at_end(document, blocks, options)


def _insert_at_index(
    document: Document,
    blocks: Sequence[ParagraphBlock],
    position: InsertPosition,
    options: InsertOptions,
) -> None:
    if position.paragraph_index is None:
        raise ValueError("paragraph_index is required when mode='index'")

    paragraphs = list(document.paragraphs)
    if position.paragraph_index <= 0:
        _insert_at_start(document, blocks, options)
        return

    if position.paragraph_index >= len(paragraphs):
        _insert_at_end(document, blocks, options)
        return

    anchor = paragraphs[position.paragraph_index]
    _insert_before_anchor(document, anchor, blocks, options)


def _insert_before_text(
    document: Document,
    blocks: Sequence[ParagraphBlock],
    position: InsertPosition,
    options: InsertOptions,
) -> None:
    if not position.anchor_text:
        raise ValueError("anchor_text is required when mode='before_text'")

    anchor = find_paragraph_by_text(document, position.anchor_text)
    _insert_before_anchor(document, anchor, blocks, options)


def _insert_after_text(
    document: Document,
    blocks: Sequence[ParagraphBlock],
    position: InsertPosition,
    options: InsertOptions,
) -> None:
    if not position.anchor_text:
        raise ValueError("anchor_text is required when mode='after_text'")

    anchor = find_paragraph_by_text(document, position.anchor_text)
    _insert_after_anchor(document, anchor, blocks, options)


def _insert_before_anchor(
    document: Document,
    anchor: Paragraph,
    blocks: Sequence[ParagraphBlock],
    options: InsertOptions,
) -> None:
    if options.page_break_before_first_block:
        page_break_paragraph = anchor.insert_paragraph_before(style=resolve_paragraph_style(document, "Normal"))
        page_break_paragraph.add_run().add_break(WD_BREAK.PAGE)

    for block in blocks:
        paragraph = anchor.insert_paragraph_before(style=resolve_paragraph_style(document, block.style))
        paragraph.add_run(block.text)


def _insert_after_anchor(
    document: Document,
    anchor: Paragraph,
    blocks: Sequence[ParagraphBlock],
    options: InsertOptions,
) -> None:
    current_anchor = anchor

    if options.page_break_before_first_block:
        current_anchor = _insert_paragraph_after(
            current_anchor,
            style=resolve_paragraph_style(document, "Normal"),
            add_page_break=True,
        )

    for block in blocks:
        current_anchor = _insert_paragraph_after(
            current_anchor,
            style=resolve_paragraph_style(document, block.style),
            text=block.text,
        )


def _insert_paragraph_after(
    paragraph: Paragraph,
    *,
    style: str,
    text: str = "",
    add_page_break: bool = False,
) -> Paragraph:
    new_paragraph_xml = OxmlElement("w:p")
    paragraph._p.addnext(new_paragraph_xml)
    new_paragraph = Paragraph(new_paragraph_xml, paragraph._parent)
    new_paragraph.style = style

    if add_page_break:
        new_paragraph.add_run().add_break(WD_BREAK.PAGE)

    if text:
        new_paragraph.add_run(text)

    return new_paragraph
